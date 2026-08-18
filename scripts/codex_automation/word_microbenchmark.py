from __future__ import annotations

from dataclasses import asdict
import json
import os
from pathlib import Path
from statistics import median
import sys
from typing import Callable
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches
from lxml import etree

from scripts.codex_automation.process import ChildResult, run_owned_child
from scripts.codex_automation.trace_profile import profile_event_trace


def _add_formatted_line(paragraph, prefix: str) -> None:
    paragraph.add_run(f"{prefix}: normal ")
    paragraph.add_run("bold").bold = True
    paragraph.add_run(" normal.")


def _set_column_spacing(path: Path, twips: int) -> None:
    temporary = path.with_suffix(".columns.tmp.docx")
    word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with ZipFile(path, "r") as source:
        members = {name: source.read(name) for name in source.namelist()}
    root = etree.fromstring(members["word/document.xml"])
    section_properties = root.find(f".//{{{word_namespace}}}sectPr")
    columns = section_properties.find(f"{{{word_namespace}}}cols") if section_properties is not None else None
    if columns is None and section_properties is not None:
        columns = etree.SubElement(section_properties, f"{{{word_namespace}}}cols")
    if columns is not None:
        columns.set(f"{{{word_namespace}}}space", str(int(twips)))
    members["word/document.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone="yes",
    )
    with ZipFile(temporary, "w", ZIP_DEFLATED) as destination:
        for name, data in members.items():
            destination.writestr(name, data)
    temporary.replace(path)


def build_text_growth_fixture(
    path: Path,
    *,
    prefix_paragraphs: int,
    measured_paragraphs: int,
    table_cells: int,
) -> Path:
    if min(prefix_paragraphs, measured_paragraphs, table_cells) < 0:
        raise ValueError("benchmark counts must be non-negative")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()

    for index in range(prefix_paragraphs):
        document.add_paragraph(f"Prefix {index:04d}: ordinary document growth text.")

    if table_cells:
        table = document.add_table(rows=table_cells, cols=1)
        for index in range(table_cells):
            _add_formatted_line(table.cell(index, 0).paragraphs[0], f"Cell {index:04d}")

    for index in range(measured_paragraphs):
        _add_formatted_line(document.add_paragraph(), f"Measured {index:04d}")

    document.save(path)
    _set_column_spacing(path, 708)
    return path


def build_table_batch_fixture(
    path: Path,
    *,
    tables: int,
    rows: int,
    columns: int,
) -> Path:
    if min(tables, rows, columns) < 0:
        raise ValueError("benchmark counts must be non-negative")
    if tables and (rows == 0 or columns == 0):
        raise ValueError("non-empty benchmark tables require rows and columns")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for table_index in range(tables):
        table = document.add_table(rows=rows, cols=columns)
        table.autofit = False
        for row_index in range(rows):
            for column_index in range(columns):
                cell = table.cell(row_index, column_index)
                cell.width = Inches(1.25 + (column_index * 0.05))
                cell.vertical_alignment = (
                    WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if (row_index + column_index) % 2
                    else WD_CELL_VERTICAL_ALIGNMENT.TOP
                )
                _add_formatted_line(
                    cell.paragraphs[0],
                    f"Table {table_index:02d} Cell {row_index:02d},{column_index:02d}",
                )
        if table_index + 1 < tables:
            document.add_paragraph(f"Table benchmark separator {table_index:02d}")
    if tables:
        document.add_paragraph("Table benchmark end")
    document.save(path)
    return path


def run_word_microbenchmark(
    source: Path,
    output_dir: Path,
    *,
    repo_root: Path,
    timeout_seconds: int = 900,
    child_executor: Callable[..., ChildResult] = run_owned_child,
    disable_table_fast_path: bool = False,
) -> dict:
    from word_replica.parser.parser import DocxParser
    from word_replica.qa.policy import run_l0_l3

    source = Path(source).resolve()
    output_dir = Path(output_dir).resolve()
    repo_root = Path(repo_root).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    child_dir = output_dir / "interactive"
    child_runner = repo_root / "scripts" / "remote_harness" / "child_runner.py"

    command = [
        sys.executable,
        str(child_runner),
        "--stage", "interactive_maximum",
        "--source", str(source),
        "--run-dir", str(child_dir),
        "--defer-l4-qa",
    ]
    if disable_table_fast_path:
        command.append("--disable-table-fast-path")
    process = child_executor(
        command,
        timeout_seconds=timeout_seconds,
        stdout_path=output_dir / "stdout.log",
        stderr_path=output_dir / "stderr.log",
        ownership_file=output_dir / "owned_word.json",
        cwd=repo_root,
        env={"PYTHONPATH": os.pathsep.join([str(repo_root / "src"), str(repo_root)])},
    )

    result_path = child_dir / "result.json"
    try:
        child_payload = json.loads(result_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        child_payload = {}

    trace_path = child_dir / "event_trace.jsonl"
    if trace_path.exists():
        performance_profile = profile_event_trace(trace_path)
    else:
        performance_profile = {"available": False, "error": "event_trace.jsonl missing"}

    output_docx = child_dir / "output.docx"
    l0_l3: dict[str, dict[str, int | bool]] = {}
    if output_docx.exists():
        parser = DocxParser()
        qa = run_l0_l3(parser.parse(source), parser.parse(output_docx))
        l0_l3 = {
            level: {
                "passed": result.passed,
                "finding_count": len(result.findings),
            }
            for level, result in qa.levels.items()
        }

    qa_passed = bool(l0_l3) and all(item["passed"] for item in l0_l3.values())
    child_status = str(child_payload.get("run_status") or "FAIL")
    status = "PASS" if not process.timed_out and process.exit_code == 0 and child_status in {"PASS", "WARN"} and qa_passed else "FAIL"
    report = {
        "status": status,
        "source": str(source),
        "output": str(output_docx) if output_docx.exists() else None,
        "reconstruction_status": child_status,
        "reconstruction_elapsed_seconds": float(child_payload.get("elapsed_seconds") or process.elapsed_seconds),
        "process": asdict(process),
        "performance_profile": performance_profile,
        "l0_l3": l0_l3,
        "reasons": list(child_payload.get("reasons") or []),
    }
    report_path = output_dir / "benchmark_report.json"
    report_path.write_text(
        json.dumps(report, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )
    return report


def run_table_batch_comparison(
    source: Path,
    output_dir: Path,
    *,
    repo_root: Path,
    runs: int = 3,
    timeout_seconds: int = 900,
    child_executor: Callable[..., ChildResult] = run_owned_child,
) -> dict:
    if int(runs) <= 0:
        raise ValueError("runs must be positive")
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    pairs = []
    for index in range(1, int(runs) + 1):
        pair_dir = output_dir / f"run-{index:02d}"
        legacy = run_word_microbenchmark(
            source,
            pair_dir / "legacy",
            repo_root=repo_root,
            timeout_seconds=timeout_seconds,
            child_executor=child_executor,
            disable_table_fast_path=True,
        )
        fast = run_word_microbenchmark(
            source,
            pair_dir / "fast",
            repo_root=repo_root,
            timeout_seconds=timeout_seconds,
            child_executor=child_executor,
            disable_table_fast_path=False,
        )
        pairs.append({"run": index, "legacy": legacy, "fast": fast})

    legacy_seconds = [
        float(pair["legacy"].get("performance_profile", {}).get("table_seconds") or 0.0)
        for pair in pairs
    ]
    fast_seconds = [
        float(pair["fast"].get("performance_profile", {}).get("table_seconds") or 0.0)
        for pair in pairs
    ]
    evaluable = all(
        pair[mode].get("status") == "PASS"
        and float(pair[mode].get("performance_profile", {}).get("table_seconds") or 0.0) > 0.0
        for pair in pairs
        for mode in ("legacy", "fast")
    )
    fidelity_equal = all(
        pair["fast"].get("l0_l3") == pair["legacy"].get("l0_l3")
        and bool(pair["fast"].get("l0_l3"))
        and all(level.get("passed") for level in pair["fast"]["l0_l3"].values())
        for pair in pairs
    )
    legacy_median = float(median(legacy_seconds))
    fast_median = float(median(fast_seconds))
    improvement = (
        100.0 * (legacy_median - fast_median) / legacy_median
        if legacy_median > 0.0
        else 0.0
    )
    report = {
        "status": "PASS" if evaluable and fidelity_equal and improvement >= 50.0 else "FAIL",
        "source": str(Path(source).resolve()),
        "runs": int(runs),
        "evaluable": evaluable,
        "fidelity_equal": fidelity_equal,
        "legacy_table_seconds": legacy_seconds,
        "fast_table_seconds": fast_seconds,
        "legacy_median_table_seconds": round(legacy_median, 6),
        "fast_median_table_seconds": round(fast_median, 6),
        "improvement_percent": round(improvement, 6),
        "pairs": pairs,
    }
    (output_dir / "comparison_report.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )
    return report
