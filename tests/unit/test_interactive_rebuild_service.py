from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from word_replica.config import InteractiveOptions, RebuildOptions
from word_replica.domain.enums import ReconstructionMode, RunStatus, VisibilityMode
from word_replica.services.audit import AuditLog
from word_replica.services.interactive_rebuild import InteractiveRebuildService
from word_replica.services.project_store import ProjectStore
from tests.fixtures.build_fixtures import build_plain_text


def test_unexpected_word_bookmarks_are_removed_from_saved_package(tmp_path):
    from lxml import etree

    output = tmp_path / "output.docx"
    document_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p>"
        b"<w:bookmarkStart w:id='1' w:name='KeepMe'/>"
        b"<w:bookmarkStart w:id='2' w:name='_GoBack'/>"
        b"<w:r><w:t>text</w:t></w:r>"
        b"<w:bookmarkEnd w:id='2'/><w:bookmarkEnd w:id='1'/>"
        b"</w:p></w:body></w:document>"
    )
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/styles.xml", b"styles")

    removed = InteractiveRebuildService._remove_unexpected_bookmarks(output, {"KeepMe"})

    assert removed == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert [node.get("{%(w)s}name" % ns) for node in root.xpath("//w:bookmarkStart", namespaces=ns)] == ["KeepMe"]


def test_unexpected_headers_are_removed_from_saved_package(tmp_path):
    from lxml import etree

    output = tmp_path / "output.docx"
    document_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' "
        b"xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'>"
        b"<w:body><w:sectPr><w:headerReference w:type='default' r:id='rId1'/></w:sectPr></w:body></w:document>"
    )
    rels_xml = (
        b"<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
        b"<Relationship Id='rId1' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/header' Target='header1.xml'/>"
        b"</Relationships>"
    )
    types_xml = (
        b"<Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'>"
        b"<Override PartName='/word/header1.xml' ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml'/>"
        b"</Types>"
    )
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/_rels/document.xml.rels", rels_xml)
        archive.writestr("[Content_Types].xml", types_xml)
        archive.writestr("word/header1.xml", b"header")

    removed = InteractiveRebuildService._remove_unexpected_headers(output, set())

    assert removed == 1
    with ZipFile(output) as archive:
        assert "word/header1.xml" not in archive.namelist()
        assert b"headerReference" not in archive.read("word/document.xml")
        assert b"header1.xml" not in archive.read("word/_rels/document.xml.rels")
        types = etree.fromstring(archive.read("[Content_Types].xml"))
        assert not types.xpath("//*[local-name()='Override' and @PartName='/word/header1.xml']")


def test_unexpected_header_shape_defaults_are_removed_when_source_omits_them(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_settings = (
        b"<w:settings xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:compat/></w:settings>"
    )
    output_settings = (
        b"<w:settings xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:hdrShapeDefaults/><w:compat/></w:settings>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/settings.xml", source_settings)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/settings.xml", output_settings)
        archive.writestr("word/unchanged.bin", b"unchanged")

    removed = InteractiveRebuildService._remove_unexpected_header_shape_defaults(output, source)

    assert removed == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/settings.xml"))
        assert archive.read("word/unchanged.bin") == b"unchanged"
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert not root.xpath("./w:hdrShapeDefaults", namespaces=ns)
    assert root.xpath("./w:compat", namespaces=ns)


def test_header_shape_defaults_are_untouched_when_source_contains_them(tmp_path):
    settings = (
        b"<w:settings xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:hdrShapeDefaults/><w:compat/></w:settings>"
    )
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    for path in (source, output):
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            archive.writestr("word/settings.xml", settings)
    before = output.read_bytes()

    removed = InteractiveRebuildService._remove_unexpected_header_shape_defaults(output, source)

    assert removed == 0
    assert output.read_bytes() == before


def test_unexpected_header_shape_defaults_are_removed_when_source_has_no_settings_part(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"document")
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "word/settings.xml",
            b"<w:settings xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
            b"<w:hdrShapeDefaults/></w:settings>",
        )

    removed = InteractiveRebuildService._remove_unexpected_header_shape_defaults(output, source)

    assert removed == 1
    with ZipFile(output) as archive:
        settings = archive.read("word/settings.xml")
    assert b"hdrShapeDefaults" not in settings


def test_template_spacing_is_removed_when_source_paragraph_has_no_properties(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:r><w:br/></w:r></w:p><w:p><w:pPr><w:jc w:val='center'/></w:pPr></w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:pPr><w:spacing w:after='200' w:line='276'/></w:pPr><w:r><w:br/></w:r></w:p>"
        b"<w:p><w:pPr><w:spacing w:after='200'/><w:jc w:val='center'/></w:pPr></w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    removed = InteractiveRebuildService._remove_template_paragraph_spacing(output, source)

    assert removed == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        paragraphs = root.xpath("//*[local-name()='body']/*[local-name()='p']")
        assert paragraphs[0].find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr") is None
        assert paragraphs[1].find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr") is not None


def test_explicit_left_alignment_is_restored_when_word_omits_default(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:pPr><w:jc w:val='left'/></w:pPr></w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:pPr><w:spacing w:after='240'/></w:pPr></w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_explicit_paragraph_alignment(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        assert root.xpath("//*[local-name()='jc' and @*[local-name()='val']='left']")


def test_explicit_paragraph_alignment_is_restored_when_word_normalizes_value(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:pPr><w:jc w:val='both'/></w:pPr></w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:pPr><w:jc w:val='left'/></w:pPr></w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_explicit_paragraph_alignment(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        assert root.xpath("//*[local-name()='jc']/@*[local-name()='val']") == ["both"]


def test_explicit_run_character_spacing_is_restored_after_word_round_trip(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:r><w:rPr><w:spacing w:val='18'/></w:rPr><w:t>text</w:t></w:r></w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:r><w:rPr><w:spacing w:val='17'/></w:rPr><w:t>text</w:t></w:r></w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_explicit_run_character_spacing(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        assert root.xpath("//*[local-name()='spacing']/@*[local-name()='val']") == ["18"]


def test_explicit_run_font_names_are_restored_only_for_safely_paired_runs(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body>"
        b"<w:p><w:r><w:rPr><w:rFonts w:ascii='Times New Roman' w:hAnsi='Times New Roman'/></w:rPr>"
        b"<w:t>Heading</w:t></w:r></w:p>"
        b"<w:p><w:r><w:rPr><w:rFonts w:ascii='Calibri'/></w:rPr><w:t>Merged </w:t></w:r>"
        b"<w:r><w:rPr><w:rFonts w:ascii='Calibri'/></w:rPr><w:t>text</w:t></w:r></w:p>"
        b"<w:p><w:r><w:rPr><w:rFonts w:ascii='Arial'/></w:rPr><w:t>Mixed </w:t></w:r>"
        b"<w:r><w:rPr><w:rFonts w:ascii='Calibri'/></w:rPr><w:t>fonts</w:t></w:r></w:p>"
        b"</w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body>"
        b"<w:p><w:r><w:t>Heading</w:t></w:r></w:p>"
        b"<w:p><w:r><w:t>Merged text</w:t></w:r></w:p>"
        b"<w:p><w:r><w:t>Mixed fonts</w:t></w:r></w:p>"
        b"</w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)
        archive.writestr("word/unchanged.bin", b"unchanged")

    restored = InteractiveRebuildService._restore_explicit_run_font_names(output, source)

    assert restored == 2
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        assert archive.read("word/unchanged.bin") == b"unchanged"
    paragraphs = root.xpath("//*[local-name()='body']/*[local-name()='p']")
    assert paragraphs[0].xpath(
        "./*[local-name()='r']/*[local-name()='rPr']/*[local-name()='rFonts']"
        "/@*[local-name()='ascii']"
    ) == ["Times New Roman"]
    assert paragraphs[0].xpath(
        "./*[local-name()='r']/*[local-name()='rPr']/*[local-name()='rFonts']"
        "/@*[local-name()='hAnsi']"
    ) == ["Times New Roman"]
    assert paragraphs[1].xpath(
        "./*[local-name()='r']/*[local-name()='rPr']/*[local-name()='rFonts']"
        "/@*[local-name()='ascii']"
    ) == ["Calibri"]
    assert not paragraphs[2].xpath(".//*[local-name()='rFonts']")


def test_explicit_run_font_names_refuse_mismatched_paragraph_topology(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body>"
        b"<w:p><w:r><w:rPr><w:rFonts w:ascii='Calibri'/></w:rPr><w:t>Repeated</w:t></w:r></w:p>"
        b"<w:p><w:r><w:rPr><w:rFonts w:ascii='Times New Roman'/></w:rPr><w:t>Repeated</w:t></w:r></w:p>"
        b"</w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:r><w:t>Repeated</w:t></w:r></w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)
    before = output.read_bytes()

    restored = InteractiveRebuildService._restore_explicit_run_font_names(output, source)

    assert restored == 0
    assert output.read_bytes() == before


def test_explicit_run_font_names_follow_run_property_schema_order(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:r><w:rPr><w:rFonts w:ascii='Times New Roman'/></w:rPr>"
        b"<w:t>Heading</w:t></w:r></w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:r><w:rPr><w:rStyle w:val='Emphasis'/><w:b/></w:rPr>"
        b"<w:t>Heading</w:t></w:r></w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_explicit_run_font_names(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    run_property_names = [
        etree.QName(child).localname
        for child in root.xpath("//*[local-name()='rPr']")[0]
    ]
    assert run_property_names == ["rStyle", "rFonts", "b"]


def test_drawing_effect_extents_are_restored_only_for_matching_geometry(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    prefix = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' "
        b"xmlns:wp='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'>"
        b"<w:body>"
    )
    source_xml = prefix + (
        b"<w:p><w:r><w:drawing><wp:inline><wp:extent cx='100' cy='200'/>"
        b"<wp:effectExtent l='0' t='0' r='0' b='0'/></wp:inline></w:drawing></w:r></w:p>"
        b"<w:p><w:r><w:drawing><wp:inline><wp:extent cx='300' cy='400'/>"
        b"<wp:effectExtent l='1' t='2' r='3' b='4'/></wp:inline></w:drawing></w:r></w:p>"
        b"</w:body></w:document>"
    )
    output_xml = prefix + (
        b"<w:p><w:r><w:drawing><wp:inline><wp:extent cx='100' cy='200'/>"
        b"<wp:effectExtent l='0' t='0' r='0' b='1270'/></wp:inline></w:drawing></w:r></w:p>"
        b"<w:p><w:r><w:drawing><wp:inline><wp:extent cx='301' cy='400'/>"
        b"<wp:effectExtent l='9' t='9' r='9' b='9'/></wp:inline></w:drawing></w:r></w:p>"
        b"</w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)
        archive.writestr("word/unchanged.bin", b"unchanged")

    restored = InteractiveRebuildService._restore_explicit_drawing_effect_extents(
        output, source
    )

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        assert archive.read("word/unchanged.bin") == b"unchanged"
    extents = root.xpath("//*[local-name()='effectExtent']")
    assert [dict(node.attrib) for node in extents] == [
        {"l": "0", "t": "0", "r": "0", "b": "0"},
        {"l": "9", "t": "9", "r": "9", "b": "9"},
    ]


def test_drawing_effect_extents_refuse_swapped_same_size_images(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    prefix = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' "
        b"xmlns:wp='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing' "
        b"xmlns:a='http://schemas.openxmlformats.org/drawingml/2006/main' "
        b"xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'>"
        b"<w:body>"
    )

    def drawing(relationship_id, bottom):
        return (
            b"<w:p><w:r><w:drawing><wp:inline><wp:extent cx='100' cy='200'/>"
            + f"<wp:effectExtent l='0' t='0' r='0' b='{bottom}'/>".encode()
            + b"<a:graphic><a:graphicData><a:blip r:embed='"
            + relationship_id.encode()
            + b"'/></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>"
        )

    source_xml = prefix + drawing("rId1", 1) + drawing("rId2", 2) + b"</w:body></w:document>"
    output_xml = prefix + drawing("rId2", 9) + drawing("rId1", 9) + b"</w:body></w:document>"
    relationships = (
        b"<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
        b"<Relationship Id='rId1' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/image' Target='media/image1.png'/>"
        b"<Relationship Id='rId2' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/image' Target='media/image2.png'/>"
        b"</Relationships>"
    )
    for path, document_xml in ((source, source_xml), (output, output_xml)):
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", document_xml)
            archive.writestr("word/_rels/document.xml.rels", relationships)
            archive.writestr("word/media/image1.png", b"first-image")
            archive.writestr("word/media/image2.png", b"second-image")
    before = output.read_bytes()

    restored = InteractiveRebuildService._restore_explicit_drawing_effect_extents(
        output, source
    )

    assert restored == 0
    assert output.read_bytes() == before


def test_drawing_effect_extents_ignore_unrestored_field_instruction_spacing(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    prefix = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' "
        b"xmlns:wp='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing' "
        b"xmlns:a='http://schemas.openxmlformats.org/drawingml/2006/main' "
        b"xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'>"
        b"<w:body><w:p><w:r><w:t>Caption</w:t></w:r><w:r><w:drawing><wp:inline>"
        b"<wp:extent cx='100' cy='200'/>"
    )
    suffix = (
        b"<a:graphic><a:graphicData><a:blip r:embed='rId1'/></a:graphicData></a:graphic>"
        b"</wp:inline></w:drawing></w:r></w:p>"
    )
    source_xml = (
        prefix
        + b"<wp:effectExtent l='0' t='0' r='0' b='0'/>"
        + suffix
        + b"<w:p><w:r><w:instrText>PAGE</w:instrText></w:r><w:r><w:t>1</w:t></w:r></w:p>"
        + b"</w:body></w:document>"
    )
    output_xml = (
        prefix
        + b"<wp:effectExtent l='0' t='0' r='0' b='1270'/>"
        + suffix
        + b"<w:p><w:r><w:instrText> PAGE </w:instrText></w:r><w:r><w:t>1</w:t></w:r></w:p>"
        + b"</w:body></w:document>"
    )
    relationships = (
        b"<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
        b"<Relationship Id='rId1' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/image' Target='media/image1.png'/>"
        b"</Relationships>"
    )
    for path, document_xml in ((source, source_xml), (output, output_xml)):
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", document_xml)
            archive.writestr("word/_rels/document.xml.rels", relationships)
            archive.writestr("word/media/image1.png", b"same-image")

    restored = InteractiveRebuildService._restore_explicit_drawing_effect_extents(
        output, source
    )

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    assert root.xpath("//*[local-name()='effectExtent']/@b") == ["0"]
    assert root.xpath("//*[local-name()='instrText']/text()") == [" PAGE "]


def test_empty_source_runs_are_restored_to_saved_package(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:r/></w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p/></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_empty_runs(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        assert len(root.xpath("//*[local-name()='p']/*[local-name()='r']")) == 1


def test_explicit_section_column_space_is_restored_after_word_rounding(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:sectPr><w:cols w:space='720'/></w:sectPr></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:sectPr><w:cols w:space='708'/></w:sectPr></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_explicit_column_space(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        assert root.xpath("//*[local-name()='cols']/@*[local-name()='space']") == ["720"]


def test_explicit_page_number_start_is_restored_after_word_save(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:sectPr><w:pgNumType w:start='1'/></w:sectPr></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:sectPr/></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_explicit_page_number_start(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        assert root.xpath("//*[local-name()='pgNumType']/@*[local-name()='start']") == ["1"]


def test_source_document_defaults_are_restored_to_saved_styles(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_styles = (
        b"<w:styles xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:docDefaults><w:pPrDefault><w:pPr><w:spacing w:line='360'/></w:pPr></w:pPrDefault></w:docDefaults></w:styles>"
    )
    output_styles = (
        b"<w:styles xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:docDefaults><w:pPrDefault><w:pPr><w:spacing w:line='240'/></w:pPr></w:pPrDefault></w:docDefaults></w:styles>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/styles.xml", source_styles)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/styles.xml", output_styles)

    restored = InteractiveRebuildService._restore_source_document_defaults(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/styles.xml"))
        assert root.xpath("//*[local-name()='spacing']/@*[local-name()='line']") == ["360"]


def test_source_theme_style_latin_fonts_are_resolved_without_changing_theme_or_cs(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_theme = b"""<a:theme xmlns:a='http://schemas.openxmlformats.org/drawingml/2006/main'>
      <a:themeElements><a:fontScheme name='Source'>
        <a:majorFont><a:latin typeface='Calibri'/><a:ea typeface='Source Major EA'/><a:cs typeface='Source Major CS'/></a:majorFont>
        <a:minorFont><a:latin typeface='Cambria'/><a:ea typeface='Source Minor EA'/><a:cs typeface='Source Minor CS'/></a:minorFont>
      </a:fontScheme><a:fmtScheme name='Source'/></a:themeElements>
    </a:theme>"""
    output_theme = b"""<a:theme xmlns:a='http://schemas.openxmlformats.org/drawingml/2006/main'>
      <a:themeElements><a:fontScheme name='Output'>
        <a:majorFont><a:latin typeface='Cambria'/><a:ea typeface='Output Major EA'/><a:cs typeface='Output Major CS'/></a:majorFont>
        <a:minorFont><a:latin typeface='Calibri'/><a:ea typeface='Output Minor EA'/><a:cs typeface='Output Minor CS'/></a:minorFont>
      </a:fontScheme><a:fmtScheme name='Output'/></a:themeElements>
    </a:theme>"""
    source_styles = b"""<w:styles xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>
      <w:style w:type='paragraph' w:styleId='Heading2'><w:name w:val='heading 2'/><w:rPr>
        <w:rFonts w:asciiTheme='majorHAnsi' w:hAnsiTheme='majorHAnsi' w:eastAsiaTheme='majorEastAsia' w:cstheme='majorBidi'/>
      </w:rPr></w:style>
      <w:style w:type='paragraph' w:styleId='Normal'><w:rPr><w:rFonts w:ascii='Times New Roman'/></w:rPr></w:style>
      <w:style w:type='paragraph' w:styleId='NoFonts'><w:rPr><w:rFonts w:asciiTheme='majorHAnsi'/></w:rPr></w:style>
      <w:style w:type='paragraph' w:styleId='SourceOnly'><w:rPr><w:rFonts w:asciiTheme='minorHAnsi'/></w:rPr></w:style>
    </w:styles>"""
    output_styles = b"""<w:styles xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>
      <w:style w:type='paragraph' w:styleId='Heading2'><w:name w:val='heading 2'/><w:rPr>
        <w:rFonts w:ascii='Times New Roman' w:hAnsi='Times New Roman' w:cs='Calibri'/>
      </w:rPr></w:style>
      <w:style w:type='paragraph' w:styleId='Normal'><w:rPr><w:rFonts w:ascii='Arial'/></w:rPr></w:style>
      <w:style w:type='paragraph' w:styleId='NoFonts'><w:rPr><w:b/></w:rPr></w:style>
    </w:styles>"""
    for path, theme, styles, theme_name in (
        (source, source_theme, source_styles, "theme2.xml"),
        (output, output_theme, output_styles, "theme7.xml"),
    ):
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            archive.writestr(f"word/theme/{theme_name}", theme)
            archive.writestr("word/styles.xml", styles)
            archive.writestr(
                "word/_rels/document.xml.rels",
                (
                    "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
                    "<Relationship Id='rIdTheme' "
                    "Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme' "
                    f"Target='theme/{theme_name}'/>"
                    "</Relationships>"
                ).encode(),
            )

    restored = InteractiveRebuildService._restore_source_theme_style_latin_fonts(
        output, source
    )

    assert restored == 1
    with ZipFile(output) as archive:
        theme = etree.fromstring(archive.read("word/theme/theme7.xml"))
        styles = etree.fromstring(archive.read("word/styles.xml"))
    assert theme.xpath(
        "string(//*[local-name()='fontScheme']/@name)"
    ) == "Output"
    assert theme.xpath(
        "string(//*[local-name()='majorFont']/*[local-name()='latin']/@typeface)"
    ) == "Cambria"
    assert theme.xpath(
        "string(//*[local-name()='minorFont']/*[local-name()='latin']/@typeface)"
    ) == "Calibri"
    assert theme.xpath(
        "string(//*[local-name()='majorFont']/*[local-name()='ea']/@typeface)"
    ) == "Output Major EA"
    assert theme.xpath(
        "string(//*[local-name()='majorFont']/*[local-name()='cs']/@typeface)"
    ) == "Output Major CS"
    assert theme.xpath(
        "string(//*[local-name()='minorFont']/*[local-name()='ea']/@typeface)"
    ) == "Output Minor EA"
    assert theme.xpath(
        "string(//*[local-name()='minorFont']/*[local-name()='cs']/@typeface)"
    ) == "Output Minor CS"
    assert theme.xpath(
        "string(//*[local-name()='fmtScheme']/@name)"
    ) == "Output"
    heading_fonts = styles.xpath(
        "//*[local-name()='style' and @*[local-name()='styleId']='Heading2']"
        "/*[local-name()='rPr']/*[local-name()='rFonts']"
    )[0]
    assert heading_fonts.xpath("string(@*[local-name()='asciiTheme'])") == ""
    assert heading_fonts.xpath("string(@*[local-name()='hAnsiTheme'])") == ""
    assert heading_fonts.xpath("string(@*[local-name()='ascii'])") == "Calibri"
    assert heading_fonts.xpath("string(@*[local-name()='hAnsi'])") == "Calibri"
    assert heading_fonts.xpath("string(@*[local-name()='cs'])") == "Calibri"
    assert heading_fonts.xpath("string(@*[local-name()='cstheme'])") == ""
    assert heading_fonts.xpath("string(@*[local-name()='eastAsiaTheme'])") == ""
    normal_fonts = styles.xpath(
        "//*[local-name()='style' and @*[local-name()='styleId']='Normal']"
        "/*[local-name()='rPr']/*[local-name()='rFonts']"
    )[0]
    assert normal_fonts.xpath("string(@*[local-name()='ascii'])") == "Arial"
    assert not styles.xpath(
        "//*[local-name()='style' and @*[local-name()='styleId']='NoFonts']"
        "/*[local-name()='rPr']/*[local-name()='rFonts']"
    )
    assert not styles.xpath(
        "//*[local-name()='style' and @*[local-name()='styleId']='SourceOnly']"
    )
    assert InteractiveRebuildService._restore_source_theme_style_latin_fonts(
        output, source
    ) == 0


def test_source_footer_parts_and_reference_types_are_restored(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    document_source = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' "
        b"xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'><w:body>"
        b"<w:sectPr><w:footerReference w:type='default' r:id='rId1'/></w:sectPr></w:body></w:document>"
    )
    document_output = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main' "
        b"xmlns:r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'><w:body>"
        b"<w:sectPr><w:footerReference w:type='even' r:id='rId1'/><w:footerReference w:type='default' r:id='rId2'/></w:sectPr>"
        b"</w:body></w:document>"
    )
    rels = b"<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'/>"
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document_source)
        archive.writestr("word/footer1.xml", b"<w:ftr xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:p><w:r><w:t>source</w:t></w:r></w:p></w:ftr>")
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document_output)
        archive.writestr("word/footer1.xml", b"output")
        archive.writestr("word/_rels/document.xml.rels", rels)

    restored = InteractiveRebuildService._restore_source_footer_stories(output, source, [{"default"}])

    assert restored == 2
    with ZipFile(output) as archive:
        assert b"source" in archive.read("word/footer1.xml")
        root = etree.fromstring(archive.read("word/document.xml"))
        assert len(root.xpath("//*[local-name()='footerReference' and @*[local-name()='type']='even']")) == 0


def test_relationship_free_source_header_story_is_restored_after_word_updates_field(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    document = b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body/></w:document>"
    source_header = (
        b"<w:hdr xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:p>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText> STYLEREF \"Heading 1\" </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>6. Results</w:t></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r></w:p></w:hdr>"
    )
    output_header = source_header.replace(b"6. Results", b"Appendices")
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr("word/header1.xml", source_header)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr("word/header1.xml", output_header)

    restored = InteractiveRebuildService._restore_relationship_free_source_headers(
        output, source, {"word/header1.xml"}
    )

    assert restored == 1
    with ZipFile(output) as archive:
        assert archive.read("word/header1.xml") == source_header


def test_source_header_with_relationships_is_not_replaced(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    document = b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body/></w:document>"
    relationships = b"<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'/>"
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr("word/header1.xml", b"source header")
        archive.writestr("word/_rels/header1.xml.rels", relationships)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr("word/header1.xml", b"output header")
        archive.writestr("word/_rels/header1.xml.rels", relationships)

    restored = InteractiveRebuildService._restore_relationship_free_source_headers(
        output, source, {"word/header1.xml"}
    )

    assert restored == 0
    with ZipFile(output) as archive:
        assert archive.read("word/header1.xml") == b"output header"


def test_nested_source_header_relationship_path_is_not_mistaken_for_relationship_free(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    document = b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body/></w:document>"
    relationships = b"<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'/>"
    header_part = "word/headers/header1.xml"
    relationship_part = "word/headers/_rels/header1.xml.rels"
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr(header_part, b"source header")
        archive.writestr(relationship_part, relationships)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr(header_part, b"output header")
        archive.writestr(relationship_part, relationships)

    restored = InteractiveRebuildService._restore_relationship_free_source_headers(
        output, source, {header_part}
    )

    assert restored == 0
    with ZipFile(output) as archive:
        assert archive.read(header_part) == b"output header"


def test_source_field_instructions_are_restored_after_word_save(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:r><w:instrText>TOC \\o \"1-3\"</w:instrText></w:r></w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        b"<w:body><w:p><w:r><w:instrText>TOC \\o \"1-3\" \\* MERGEFORMAT</w:instrText></w:r></w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_source_field_instructions(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        assert root.xpath("string(//*[local-name()='instrText'])") == 'TOC \\o "1-3"'


def test_word_simple_field_is_restored_to_source_complex_field_structure(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r>"
        b"<w:r><w:instrText xml:space='preserve'> REF ref_tab_1 \\h </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r>"
        b"<w:r><w:t>SOURCE RESULT</w:t></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"</w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:fldSimple w:instr=' REF ref_tab_1 \\h \\* MERGEFORMAT '>"
        b"<w:r><w:rPr><w:b/></w:rPr><w:t>OUTPUT RESULT</w:t></w:r>"
        b"</w:fldSimple></w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_source_field_instructions(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert not root.xpath("//w:fldSimple", namespaces=ns)
    assert [
        node.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType")
        for node in root.xpath("//w:fldChar", namespaces=ns)
    ] == ["begin", "separate", "end"]
    assert root.xpath("string(//w:instrText)", namespaces=ns) == " REF ref_tab_1 \\h "
    assert root.xpath("string(//w:fldChar[@w:fldCharType='separate']/following::w:t[1])", namespaces=ns) == "OUTPUT RESULT"
    assert root.xpath("boolean(//w:t[text()='OUTPUT RESULT']/../w:rPr/w:b)", namespaces=ns)


def test_nested_word_simple_fields_are_both_restored_to_complex_structure(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r>"
        b"<w:r><w:instrText> IF OUTER </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r>"
        b"<w:r><w:instrText> REF INNER </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r>"
        b"<w:r><w:t>SOURCE INNER</w:t></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"</w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:fldSimple w:instr=' IF OUTER \\* MERGEFORMAT '>"
        b"<w:fldSimple w:instr=' REF INNER \\* MERGEFORMAT '>"
        b"<w:r><w:t>OUTPUT INNER</w:t></w:r>"
        b"</w:fldSimple></w:fldSimple>"
        b"</w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_source_field_instructions(output, source)

    assert restored == 2
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert not root.xpath("//w:fldSimple", namespaces=ns)
    assert len(root.xpath("//w:fldChar[@w:fldCharType='begin']", namespaces=ns)) == 2
    assert root.xpath("string(//w:t[text()='OUTPUT INNER'])", namespaces=ns) == "OUTPUT INNER"


def test_split_field_instructions_keep_field_alignment_and_source_run_structure(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r>"
        b"<w:r><w:instrText> REF </w:instrText></w:r><w:r><w:instrText>first </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>ONE</w:t></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r>"
        b"<w:r><w:instrText> REF second </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>TWO</w:t></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"</w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r>"
        b"<w:r><w:instrText> REF first \\* MERGEFORMAT </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>OUTPUT ONE</w:t></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r>"
        b"<w:r><w:instrText> REF second \\* MERGEFORMAT </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>OUTPUT TWO</w:t></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"</w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_source_field_instructions(output, source)

    assert restored == 2
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert root.xpath("//w:instrText/text()", namespaces=ns) == [" REF ", "first ", " REF second "]
    assert root.xpath("//w:t/text()", namespaces=ns) == ["OUTPUT ONE", "OUTPUT TWO"]


def test_simple_field_expansion_preserves_source_field_flags_run_formatting_and_output_tail(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:r><w:rPr><w:b/></w:rPr><w:fldChar w:fldCharType='begin' w:fldLock='true' w:dirty='true'/></w:r>"
        b"<w:r><w:rPr><w:i/></w:rPr><w:instrText> REF locked </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r>"
        b"<w:r><w:t>SOURCE</w:t></w:r><w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"</w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:fldSimple w:instr=' REF locked \\* MERGEFORMAT '><w:r><w:t>OUTPUT</w:t></w:r></w:fldSimple>TAIL"
        b"</w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_source_field_instructions(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    begin = root.xpath("//w:fldChar[@w:fldCharType='begin']", namespaces=ns)[0]
    assert begin.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldLock") == "true"
    assert begin.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}dirty") == "true"
    assert begin.xpath("boolean(../w:rPr/w:b)", namespaces=ns)
    assert root.xpath("boolean(//w:instrText/../w:rPr/w:i)", namespaces=ns)
    assert root.xpath("string(//w:t)", namespaces=ns) == "OUTPUT"
    end_run = root.xpath("//w:fldChar[@w:fldCharType='end']/..", namespaces=ns)[0]
    assert end_run.tail == "TAIL"


def test_field_restoration_rejects_unmatched_field_topology(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText> REF one </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>ONE</w:t></w:r><w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText> REF two </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>TWO</w:t></w:r><w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"</w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:fldSimple w:instr=' REF one \\* MERGEFORMAT '><w:r><w:t>OUTPUT</w:t></w:r></w:fldSimple>"
        b"</w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_source_field_instructions(output, source)

    assert restored == 0
    with ZipFile(output) as archive:
        assert archive.read("word/document.xml") == output_xml


def test_cross_paragraph_field_does_not_block_restoration_of_other_fields(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body>"
        b"<w:p><w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText> TOC </w:instrText></w:r></w:p>"
        b"<w:p><w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>SOURCE TOC</w:t></w:r></w:p>"
        b"<w:p><w:r><w:fldChar w:fldCharType='end'/></w:r></w:p>"
        b"<w:p><w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText> REF flat </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>SOURCE FLAT</w:t></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r></w:p>"
        b"</w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body>"
        b"<w:p><w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText> TOC \\* MERGEFORMAT </w:instrText></w:r></w:p>"
        b"<w:p><w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>OUTPUT TOC</w:t></w:r></w:p>"
        b"<w:p><w:r><w:fldChar w:fldCharType='end'/></w:r></w:p>"
        b"<w:p><w:fldSimple w:instr=' REF flat \\* MERGEFORMAT '><w:r><w:t>OUTPUT FLAT</w:t></w:r></w:fldSimple></w:p>"
        b"</w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_source_field_instructions(output, source)

    assert restored == 2
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert root.xpath("//w:instrText/text()", namespaces=ns) == [" TOC ", " REF flat "]
    assert not root.xpath("//w:fldSimple", namespaces=ns)
    assert root.xpath("//w:t/text()", namespaces=ns) == ["OUTPUT TOC", "OUTPUT FLAT"]


def test_outer_field_does_not_overwrite_nested_output_result_inside_instruction(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText> IF </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText> REF nested </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>SOURCE INNER</w:t></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r><w:r><w:instrText> = 1 </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>SOURCE OUTER</w:t></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"</w:p></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:p>"
        b"<w:r><w:fldChar w:fldCharType='begin'/></w:r><w:r><w:instrText> IF \\* MERGEFORMAT </w:instrText></w:r>"
        b"<w:fldSimple w:instr=' REF nested \\* MERGEFORMAT '><w:r><w:t>OUTPUT INNER</w:t></w:r></w:fldSimple>"
        b"<w:r><w:instrText> = 1 </w:instrText></w:r>"
        b"<w:r><w:fldChar w:fldCharType='separate'/></w:r><w:r><w:t>OUTPUT OUTER</w:t></w:r>"
        b"<w:r><w:fldChar w:fldCharType='end'/></w:r>"
        b"</w:p></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_source_field_instructions(output, source)

    assert restored == 2
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert not root.xpath("//w:fldSimple", namespaces=ns)
    assert root.xpath("//w:t/text()", namespaces=ns) == ["OUTPUT INNER", "OUTPUT OUTER"]
    assert root.xpath("//w:instrText/text()", namespaces=ns) == [" IF ", " REF nested ", " = 1 "]


def test_source_autofit_table_layout_is_restored(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body>"
        b"<w:tbl><w:tblPr/><w:tr><w:tc><w:tcPr/></w:tc></w:tr></w:tbl></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body>"
        b"<w:tbl><w:tblPr><w:tblLayout w:type='fixed'/></w:tblPr><w:tblGrid><w:gridCol w:w='1000'/></w:tblGrid>"
        b"<w:tr><w:tc><w:tcPr><w:tcW w:w='1000' w:type='dxa'/></w:tcPr></w:tc></w:tr></w:tbl></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    removed = InteractiveRebuildService._restore_source_table_layout(output, source)

    assert removed == 3
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        assert not root.xpath("//*[local-name()='tblGrid']|//*[local-name()='tblLayout']|//*[local-name()='tcW']")


def test_source_table_property_topology_is_restored_after_word_materializes_defaults(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:tbl>"
        b"<w:tblPr><w:tblW w:w='0' w:type='auto'/><w:tblBorders>"
        b"<w:top w:val='single' w:sz='12' w:color='1B2F4B'/><w:left w:val='nil'/>"
        b"</w:tblBorders></w:tblPr><w:tblGrid><w:gridCol w:w='1000'/></w:tblGrid>"
        b"<w:tr><w:trPr><w:cantSplit/></w:trPr><w:tc><w:tcPr>"
        b"<w:tcW w:w='1000' w:type='dxa'/><w:shd w:fill='F2F5FA'/><w:tcBorders>"
        b"<w:bottom w:val='single' w:sz='8'/></w:tcBorders></w:tcPr>"
        b"<w:p><w:r><w:t>SOURCE CONTENT</w:t></w:r></w:p></w:tc></w:tr>"
        b"</w:tbl></w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body><w:tbl>"
        b"<w:tblPr><w:tblW w:w='1000' w:type='dxa'/><w:tblBorders>"
        b"<w:top w:val='single' w:sz='12' w:color='1B2F4B'/>"
        b"</w:tblBorders><w:shd w:fill='auto'/></w:tblPr><w:tblGrid><w:gridCol w:w='1200'/></w:tblGrid>"
        b"<w:tr><w:trPr><w:tblHeader/></w:trPr><w:tc><w:tcPr>"
        b"<w:tcW w:w='1200' w:type='dxa'/><w:shd w:fill='auto'/><w:tcBorders>"
        b"<w:top w:val='single' w:sz='12' w:color='1B2F4B'/><w:bottom w:val='nil'/>"
        b"</w:tcBorders></w:tcPr><w:p><w:r><w:t>OUTPUT CONTENT</w:t></w:r></w:p></w:tc></w:tr>"
        b"</w:tbl></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_source_table_layout(output, source)

    assert restored == 4
    with ZipFile(source) as archive:
        source_root = etree.fromstring(archive.read("word/document.xml"))
    with ZipFile(output) as archive:
        output_root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    for path in ("//w:tblPr", "//w:tblGrid", "//w:trPr", "//w:tcPr"):
        assert etree.tostring(output_root.xpath(path, namespaces=ns)[0]) == etree.tostring(
            source_root.xpath(path, namespaces=ns)[0]
        )
    assert output_root.xpath("string(//w:t)", namespaces=ns) == "OUTPUT CONTENT"


def test_source_table_properties_are_not_restored_when_nested_topology_differs(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body>"
        b"<w:tbl><w:tblPr><w:tblW w:w='1000' w:type='dxa'/></w:tblPr><w:tr><w:tc>"
        b"<w:p><w:r><w:t>SOURCE</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
        b"</w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body>"
        b"<w:tbl><w:tblPr><w:tblW w:w='2000' w:type='dxa'/></w:tblPr><w:tr><w:tc>"
        b"<w:p><w:r><w:t>OUTPUT</w:t></w:r></w:p>"
        b"<w:tbl><w:tblPr/><w:tr><w:tc><w:p/></w:tc></w:tr></w:tbl>"
        b"</w:tc></w:tr></w:tbl></w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    before = output.read_bytes()
    restored = InteractiveRebuildService._restore_source_table_layout(output, source)

    assert restored == 0
    assert output.read_bytes() == before


def test_source_row_properties_are_inserted_after_table_property_exceptions(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    source_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body>"
        b"<w:tbl><w:tr><w:tblPrEx/><w:trPr><w:cantSplit/></w:trPr><w:tc><w:p/></w:tc></w:tr></w:tbl>"
        b"</w:body></w:document>"
    )
    output_xml = (
        b"<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'><w:body>"
        b"<w:tbl><w:tr><w:tblPrEx/><w:tc><w:p/></w:tc></w:tr></w:tbl>"
        b"</w:body></w:document>"
    )
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", source_xml)
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", output_xml)

    restored = InteractiveRebuildService._restore_source_table_layout(output, source)

    assert restored == 1
    with ZipFile(output) as archive:
        output_root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    row = output_root.xpath("//w:tr", namespaces=ns)[0]
    assert [etree.QName(child).localname for child in row[:2]] == ["tblPrEx", "trPr"]


def test_blocked_preflight_never_creates_word_controller(tmp_path):
    source=build_plain_text(tmp_path/"source.docx")
    store=ProjectStore(tmp_path/"projects")
    options=RebuildOptions(reconstruction_mode=ReconstructionMode.INTERACTIVE)
    paths=store.create_project(source, options)
    audit=AuditLog(paths.logs_dir/"audit.jsonl")
    service=InteractiveRebuildService(word_probe=lambda:False)
    prepared=service.prepare(source, options, paths, audit)
    called=[]
    result=service.start(prepared, controller_factory=lambda: called.append(True))
    assert called == []
    assert result.status is RunStatus.FAIL
    assert any("Word" in reason for reason in result.reasons)
    assert (paths.working_dir/"blueprint.json").exists()
    assert (paths.logs_dir/"preflight.json").exists()


def test_ready_start_executes_blueprint_and_writes_truthful_checkpoint(tmp_path):
    from word_replica.interactive.control import InteractiveRunControl
    from word_replica.qa.render import RenderQaResult
    source=build_plain_text(tmp_path/"source.docx")
    store=ProjectStore(tmp_path/"projects")
    options=RebuildOptions(reconstruction_mode=ReconstructionMode.INTERACTIVE,
        interactive=InteractiveOptions(checkpoint_event_interval=1))
    paths=store.create_project(source, options)
    audit=AuditLog(paths.logs_dir/"audit.jsonl")
    def exporter(docx, pdf, visible=False):
        Path(pdf).write_bytes(b"pdf")
        return Path(pdf)
    def comparer(source_pdf, rebuilt_pdf, qa_dir):
        return RenderQaResult(True, True, True, 1, 1, [], [])
    service=InteractiveRebuildService(
        word_probe=lambda:True, project_store=store,
        pdf_exporter=exporter, pdf_comparer=comparer,
    )
    prepared=service.prepare(source, options, paths, audit)

    class FakeController:
        def __init__(self): self.opened=False; self.events=[]; self.range=0
        def open_blank(self): self.opened=True
        def set_asset_resolver(self,resolver): self.resolver=resolver
        def execute_event(self,event):
            self.events.append(event.event_type)
            if event.event_type == "InsertCharacter": self.range += 1
            if event.event_type == "InsertText": self.range += len(event.payload["text"])
        def save(self,path):
            import shutil
            shutil.copy2(source, path)
        def current_state_snapshot(self): return {"story":"body","range_start":self.range,"range_end":self.range,"paragraph_started":True}
        def set_custom_property(self,name,value): pass
        def close(self): pass
    controller=FakeController()
    control=InteractiveRunControl(); control.start()
    result=service.start(prepared, controller_factory=lambda:controller, control=control)
    assert controller.opened is True
    assert result.output_path and result.output_path.exists()
    assert result.save_count > 0
    assert (paths.logs_dir/"interactive_checkpoint.json").exists()
    assert store.get_project(paths.project_id)["status"] in {"VERIFYING","COMPLETED"}


def test_default_controller_receives_background_visibility(tmp_path, monkeypatch):
    import shutil
    from word_replica.qa.render import RenderQaResult
    from word_replica.services import interactive_rebuild as service_module

    source = build_plain_text(tmp_path / "source.docx")
    store = ProjectStore(tmp_path / "projects")
    options = RebuildOptions(
        reconstruction_mode=ReconstructionMode.INTERACTIVE,
        visibility=VisibilityMode.BACKGROUND,
        interactive=InteractiveOptions(checkpoint_event_interval=9999),
    )
    paths = store.create_project(source, options)
    audit = AuditLog(paths.logs_dir / "audit.jsonl")
    seen_visibility = []

    class FakeController:
        def __init__(self, *, visible):
            seen_visibility.append(visible)
            self.position = 0
        def open_blank(self): pass
        def set_asset_resolver(self, resolver): self.resolver = resolver
        def execute_event(self, event):
            if event.event_type == "InsertText":
                self.position += len(event.payload["text"])
        def save(self, path): shutil.copy2(source, path)
        def current_state_snapshot(self):
            return {"story": "body", "range_start": self.position, "range_end": self.position,
                    "paragraph_started": True}
        def set_custom_property(self, name, value): pass
        def close(self): pass

    monkeypatch.setattr(service_module, "InteractiveWordController", FakeController)
    def exporter(docx, pdf, visible=False):
        Path(pdf).write_bytes(b"pdf")
        return Path(pdf)

    service = InteractiveRebuildService(
        word_probe=lambda: True,
        project_store=store,
        pdf_exporter=exporter,
        pdf_comparer=lambda source_pdf, rebuilt_pdf, qa_dir: RenderQaResult(
            True, True, True, 1, 1, [], []
        ),
    )
    prepared = service.prepare(source, options, paths, audit)

    service.start(prepared)

    assert seen_visibility == [False]


def test_resume_rejects_changed_source_before_opening_word(tmp_path):
    from word_replica.interactive.checkpoints import InteractiveCheckpoint, InteractiveCheckpointStore
    from word_replica.interactive.control import InteractiveRunControl
    from word_replica.interactive.blueprint import BlueprintCompiler
    from word_replica.parser.parser import DocxParser
    from word_replica.services.source_guard import sha256_file

    source=build_plain_text(tmp_path/"source.docx")
    store=ProjectStore(tmp_path/"projects")
    options=RebuildOptions(reconstruction_mode=ReconstructionMode.INTERACTIVE)
    paths=store.create_project(source, options)
    model=DocxParser().parse(source); bp=BlueprintCompiler().compile(model)
    output=paths.output_dir/"partial.docx"; output.write_bytes(b"partial")
    cp=InteractiveCheckpoint.now(project_id=paths.project_id, source_sha256=sha256_file(source),
        source_model_fingerprint=bp.source_model_fingerprint, blueprint_fingerprint=bp.fingerprint,
        blueprint_schema_version=bp.schema_version, output_path=str(output), output_sha256=sha256_file(output),
        last_completed_event_index=0, section_element_id=None, block_element_id=None, table_element_id=None,
        cell_element_id=None, save_sequence=1, settings={"speed_mode":"fast","characters_per_second":25,
        "object_step_delay_ms":0,"fidelity":"maximum","checkpoint_after_tables":True,
        "checkpoint_after_images":True,"checkpoint_after_sections":True,"checkpoint_event_interval":500,
        "verify_during_run":True,"block_on_unsupported":True,"allow_preserved_objects":False}, status="STOPPED",
        story="body", range_start=1, range_end=1, paragraph_started=True)
    InteractiveCheckpointStore(paths.logs_dir/"interactive_checkpoint.json").write(cp)
    store.set_status(paths.project_id,"STOPPED")
    source.write_bytes(b"changed")
    called=[]
    service=InteractiveRebuildService(project_store=store, word_probe=lambda:True)
    control=InteractiveRunControl(); control.start()
    result=service.resume(paths.project_id, control, controller_factory=lambda:called.append(True))
    assert result.status is RunStatus.FAIL
    assert called == []
    assert any("source hash" in r.lower() for r in result.reasons)


def test_resume_opens_partial_restores_range_and_continues_after_checkpoint(tmp_path):
    from dataclasses import asdict
    from word_replica.interactive.blueprint import BlueprintCompiler
    from word_replica.interactive.checkpoints import InteractiveCheckpoint, InteractiveCheckpointStore
    from word_replica.interactive.control import InteractiveRunControl
    from word_replica.parser.parser import DocxParser
    from word_replica.services.source_guard import sha256_file

    source=build_plain_text(tmp_path/"source.docx")
    store=ProjectStore(tmp_path/"projects")
    options=RebuildOptions(reconstruction_mode=ReconstructionMode.INTERACTIVE,
        interactive=InteractiveOptions(checkpoint_event_interval=999))
    paths=store.create_project(source, options)
    model=DocxParser().parse(source); bp=BlueprintCompiler().compile(model)
    first_char = next(i for i,e in enumerate(bp.events) if e.event_type == "InsertText")
    output=paths.output_dir/"partial.docx"; output.write_bytes(b"partial")
    settings=asdict(options.interactive)
    settings["speed_mode"]=options.interactive.speed_mode.value; settings["fidelity"]=options.interactive.fidelity.value
    cp=InteractiveCheckpoint.now(project_id=paths.project_id, source_sha256=sha256_file(source),
        source_model_fingerprint=bp.source_model_fingerprint, blueprint_fingerprint=bp.fingerprint,
        blueprint_schema_version=bp.schema_version, output_path=str(output), output_sha256=sha256_file(output),
        last_completed_event_index=first_char, section_element_id=None, block_element_id=None, table_element_id=None,
        cell_element_id=None, save_sequence=1, settings=settings, status="STOPPED", story="body",
        range_start=7, range_end=7, paragraph_started=True)
    InteractiveCheckpointStore(paths.logs_dir/"interactive_checkpoint.json").write(cp)
    store.set_status(paths.project_id,"STOPPED")

    class FakeController:
        def __init__(self): self.opened=None; self.restored=None; self.events=[]; self.range=7
        def open_existing(self,path): self.opened=Path(path)
        def restore_checkpoint_state(self,checkpoint): self.restored=(checkpoint.range_start, checkpoint.range_end)
        def set_asset_resolver(self,resolver): self.resolver=resolver
        def execute_event(self,event): self.events.append(event)
        def save(self,path): Path(path).write_bytes(b"resumed")
        def current_state_snapshot(self): return {"story":"body","range_start":8,"range_end":8,"paragraph_started":True}
        def set_custom_property(self,name,value): pass
        def close(self): pass
    controller=FakeController()
    control=InteractiveRunControl(); control.start()
    service=InteractiveRebuildService(project_store=store, word_probe=lambda:True)
    result=service.resume(paths.project_id, control, controller_factory=lambda:controller)
    assert controller.opened == output
    assert controller.restored == (7,7)
    assert controller.events[0] == bp.events[first_char+1]
    assert result.output_path == output
    assert result.save_count >= 2


def test_completed_interactive_run_requires_final_l0_l3_and_l4_before_pass(tmp_path):
    import shutil
    from word_replica.interactive.control import InteractiveRunControl
    from word_replica.qa.render import RenderQaResult
    source=build_plain_text(tmp_path/"source.docx")
    store=ProjectStore(tmp_path/"projects")
    options=RebuildOptions(reconstruction_mode=ReconstructionMode.INTERACTIVE,
        interactive=InteractiveOptions(checkpoint_event_interval=9999))
    paths=store.create_project(source,options); audit=AuditLog(paths.logs_dir/"audit.jsonl")
    exports=[]
    def exporter(docx,pdf,visible=False):
        exports.append(Path(docx)); Path(pdf).write_bytes(b"pdf"); return Path(pdf)
    def comparer(source_pdf,rebuilt_pdf,qa_dir):
        return RenderQaResult(True,True,True,1,1,[],[])
    service=InteractiveRebuildService(project_store=store, word_probe=lambda:True,
        pdf_exporter=exporter, pdf_comparer=comparer)
    prepared=service.prepare(source,options,paths,audit)
    class Controller:
        def __init__(self): self.props={}; self.pos=0
        def open_blank(self): pass
        def set_asset_resolver(self,resolver): pass
        def execute_event(self,event):
            if event.event_type=="InsertCharacter": self.pos+=1
            if event.event_type=="InsertText": self.pos+=len(event.payload["text"])
        def current_state_snapshot(self): return {"document_identity":"d","story":"body","range_start":self.pos,"range_end":self.pos,"section_index":0,"table_element_id":None,"cell_element_id":None}
        def save(self,path): shutil.copy2(source,path)
        def set_custom_property(self,name,value): self.props[name]=value
        def close(self): pass
    controller=Controller(); control=InteractiveRunControl(); control.start()
    result=service.start(prepared,controller_factory=lambda:controller,control=control)
    assert result.status is RunStatus.PASS
    assert result.qa_report_path and result.qa_report_path.exists()
    assert exports == [source.resolve(), result.output_path]
    assert store.get_project(paths.project_id)["status"] == "COMPLETED"
    assert int(controller.props["WordReplicaActualSaveCount"]) == result.save_count
    audit_rows = [__import__("json").loads(line) for line in (paths.logs_dir/"audit.jsonl").read_text(encoding="utf-8").splitlines() if line]
    metrics = [row for row in audit_rows if row.get("event_type") == "INTERACTIVE_METRICS"]
    assert metrics
    assert metrics[-1]["payload"]["completed_characters"] == prepared.blueprint.total_visible_characters
    assert metrics[-1]["payload"]["completed_events"] == prepared.blueprint.total_events


def test_deferred_l4_still_seals_output_and_runs_l0_l3_without_pdf_export(tmp_path, monkeypatch):
    import json
    import shutil
    from word_replica.services.checkpoints import CheckpointManager

    source = build_plain_text(tmp_path / "source.docx")
    store = ProjectStore(tmp_path / "projects")
    options = RebuildOptions(reconstruction_mode=ReconstructionMode.INTERACTIVE)
    paths = store.create_project(source, options)
    audit = AuditLog(paths.logs_dir / "audit.jsonl")
    service = InteractiveRebuildService(
        project_store=store,
        word_probe=lambda: True,
        run_l4_qa=False,
        pdf_exporter=lambda *args, **kwargs: (_ for _ in ()).throw(
            AssertionError("deferred L4 must not export PDFs")
        ),
    )
    restored_effect_extents = []

    def restore_effect_extents(output_path, source_path):
        restored_effect_extents.append((Path(output_path), Path(source_path)))
        return 2

    monkeypatch.setattr(
        service,
        "_restore_explicit_drawing_effect_extents",
        restore_effect_extents,
    )
    prepared = service.prepare(source, options, paths, audit)
    output = paths.output_dir / "reconstructed.docx"
    save_manager = CheckpointManager(paths.logs_dir / "save_history.jsonl", audit)

    class Renderer:
        def __init__(self):
            self.properties = {}

        def set_custom_property(self, name, value):
            self.properties[name] = value

        def save(self, path):
            shutil.copy2(source, path)

        def current_state_snapshot(self):
            return {"story": "body", "range_start": 0, "range_end": 0}

    renderer = Renderer()
    result = service._finalize_qa(prepared, output, save_manager, audit, renderer)

    assert result.status is RunStatus.PASS
    assert result.output_path == output
    assert result.qa_report_path and result.qa_report_path.exists()
    assert result.save_count == 1
    assert renderer.properties["WordReplicaActualSaveCount"] == 1
    assert any(warning.code == "L4_DEFERRED" for warning in result.warnings)
    assert restored_effect_extents == [(output, source)]
    audit_rows = [
        json.loads(line)
        for line in (paths.logs_dir / "audit.jsonl").read_text(encoding="utf-8").splitlines()
        if line
    ]
    events = [
        row for row in audit_rows
        if row.get("event_type") == "EXPLICIT_DRAWING_EFFECT_EXTENTS_RESTORED"
    ]
    assert events[-1]["payload"] == {"count": 2}


def test_finalization_removes_unexpected_header_shape_defaults_and_audits_it(tmp_path):
    import json
    import shutil

    from lxml import etree
    from word_replica.services.checkpoints import CheckpointManager

    source = build_plain_text(tmp_path / "source.docx")
    store = ProjectStore(tmp_path / "projects")
    options = RebuildOptions(reconstruction_mode=ReconstructionMode.INTERACTIVE)
    paths = store.create_project(source, options)
    audit = AuditLog(paths.logs_dir / "audit.jsonl")
    service = InteractiveRebuildService(
        project_store=store,
        word_probe=lambda: True,
        run_l4_qa=False,
    )
    prepared = service.prepare(source, options, paths, audit)
    output = paths.output_dir / "reconstructed.docx"
    save_manager = CheckpointManager(paths.logs_dir / "save_history.jsonl", audit)

    class Renderer:
        def set_custom_property(self, name, value):
            pass

        def save(self, path):
            shutil.copy2(source, path)
            with ZipFile(path) as archive:
                parts = {name: archive.read(name) for name in archive.namelist()}
            settings = etree.fromstring(parts["word/settings.xml"])
            namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            assert not settings.xpath(
                "./w:hdrShapeDefaults", namespaces={"w": namespace}
            )
            settings.insert(0, etree.Element(f"{{{namespace}}}hdrShapeDefaults"))
            parts["word/settings.xml"] = etree.tostring(
                settings, xml_declaration=True, encoding="UTF-8", standalone=True
            )
            mutated = tmp_path / "mutated.docx"
            with ZipFile(mutated, "w", ZIP_DEFLATED) as archive:
                for part_name, data in parts.items():
                    archive.writestr(part_name, data)
            mutated.replace(path)

        def current_state_snapshot(self):
            return {"story": "body", "range_start": 0, "range_end": 0}

        def close(self):
            pass

    result = service._finalize_qa(
        prepared,
        output,
        save_manager,
        audit,
        Renderer(),
    )

    assert result.status is RunStatus.PASS
    with ZipFile(output) as archive:
        settings = etree.fromstring(archive.read("word/settings.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert not settings.xpath("./w:hdrShapeDefaults", namespaces=ns)
    audit_rows = [
        json.loads(line)
        for line in (paths.logs_dir / "audit.jsonl").read_text(encoding="utf-8").splitlines()
        if line
    ]
    events = [
        row for row in audit_rows
        if row.get("event_type") == "UNEXPECTED_HEADER_SHAPE_DEFAULTS_REMOVED"
    ]
    assert events[-1]["payload"] == {"count": 1}


def test_finalization_restores_explicit_run_font_names_and_audits_it(tmp_path):
    import json
    import shutil

    from lxml import etree
    from word_replica.services.checkpoints import CheckpointManager

    source = build_plain_text(tmp_path / "source.docx")
    with ZipFile(source) as archive:
        source_parts = {name: archive.read(name) for name in archive.namelist()}
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ns = {"w": namespace}
    source_root = etree.fromstring(source_parts["word/document.xml"])
    source_run = source_root.xpath("//w:body//w:r[w:t][1]", namespaces=ns)[0]
    source_rpr = etree.Element(f"{{{namespace}}}rPr")
    source_fonts = etree.SubElement(source_rpr, f"{{{namespace}}}rFonts")
    source_fonts.set(f"{{{namespace}}}ascii", "Times New Roman")
    source_fonts.set(f"{{{namespace}}}hAnsi", "Times New Roman")
    source_run.insert(0, source_rpr)
    source_parts["word/document.xml"] = etree.tostring(
        source_root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    rewritten_source = tmp_path / "source-with-fonts.docx"
    with ZipFile(rewritten_source, "w", ZIP_DEFLATED) as archive:
        for part_name, data in source_parts.items():
            archive.writestr(part_name, data)
    rewritten_source.replace(source)

    store = ProjectStore(tmp_path / "projects")
    options = RebuildOptions(reconstruction_mode=ReconstructionMode.INTERACTIVE)
    paths = store.create_project(source, options)
    audit = AuditLog(paths.logs_dir / "audit.jsonl")
    service = InteractiveRebuildService(
        project_store=store,
        word_probe=lambda: True,
        run_l4_qa=False,
    )
    prepared = service.prepare(source, options, paths, audit)
    output = paths.output_dir / "reconstructed.docx"
    save_manager = CheckpointManager(paths.logs_dir / "save_history.jsonl", audit)

    class Renderer:
        def set_custom_property(self, name, value):
            pass

        def save(self, path):
            shutil.copy2(source, path)
            with ZipFile(path) as archive:
                parts = {name: archive.read(name) for name in archive.namelist()}
            root = etree.fromstring(parts["word/document.xml"])
            run = root.xpath("//w:body//w:r[w:t][1]", namespaces=ns)[0]
            run.remove(run.find(f"{{{namespace}}}rPr"))
            parts["word/document.xml"] = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )
            mutated = tmp_path / "mutated-fonts.docx"
            with ZipFile(mutated, "w", ZIP_DEFLATED) as archive:
                for part_name, data in parts.items():
                    archive.writestr(part_name, data)
            mutated.replace(path)

        def current_state_snapshot(self):
            return {"story": "body", "range_start": 0, "range_end": 0}

        def close(self):
            pass

    result = service._finalize_qa(
        prepared,
        output,
        save_manager,
        audit,
        Renderer(),
    )

    assert result.status is RunStatus.PASS
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    restored_fonts = root.xpath(
        "//w:body//w:r[w:t][1]/w:rPr/w:rFonts", namespaces=ns
    )
    assert len(restored_fonts) == 1
    assert restored_fonts[0].get(f"{{{namespace}}}ascii") == "Times New Roman"
    audit_rows = [
        json.loads(line)
        for line in (paths.logs_dir / "audit.jsonl").read_text(encoding="utf-8").splitlines()
        if line
    ]
    events = [
        row for row in audit_rows
        if row.get("event_type") == "EXPLICIT_RUN_FONT_NAMES_RESTORED"
    ]
    assert events[-1]["payload"] == {"count": 1}


def test_live_verification_mismatch_pauses_at_saved_table_boundary(tmp_path):
    import shutil
    from word_replica.domain.reconstruction import LiveVerificationResult
    from word_replica.interactive.control import InteractiveRunControl
    from tests.fixtures.build_fixtures import build_tables_merged

    source = build_tables_merged(tmp_path / "table.docx")
    store = ProjectStore(tmp_path / "projects")
    options = RebuildOptions(
        reconstruction_mode=ReconstructionMode.INTERACTIVE,
        interactive=InteractiveOptions(
            verify_during_run=True,
            checkpoint_after_tables=True,
            checkpoint_event_interval=99999,
        ),
    )
    paths = store.create_project(source, options)
    audit = AuditLog(paths.logs_dir / "audit.jsonl")

    class FailingVerifier:
        def __init__(self): self.calls = []
        def verify_boundary(self, source_model, output_path, boundary):
            self.calls.append((Path(output_path), dict(boundary)))
            return LiveVerificationResult(
                "PAUSED_FIDELITY_MISMATCH", str(boundary["name"]), False, ("table count mismatch",)
            )

    verifier = FailingVerifier()
    service = InteractiveRebuildService(
        word_probe=lambda: True,
        project_store=store,
        live_verifier=verifier,
    )
    prepared = service.prepare(source, options, paths, audit)

    class Controller:
        def __init__(self): self.events=[]; self.pos=0
        def open_blank(self): pass
        def set_asset_resolver(self, resolver): pass
        def execute_event(self, event):
            self.events.append(event.event_type)
            if event.event_type == "InsertCharacter": self.pos += 1
            if event.event_type == "InsertText": self.pos += len(event.payload["text"])
        def current_state_snapshot(self):
            return {"document_identity":"d","story":"body","range_start":self.pos,"range_end":self.pos,
                    "section_index":0,"table_element_id":None,"cell_element_id":None}
        def save(self, path): shutil.copy2(source, path)
        def set_custom_property(self, name, value): pass
        def close(self): pass

    controller=Controller(); control=InteractiveRunControl(); control.start()
    result=service.start(prepared, controller_factory=lambda:controller, control=control)
    assert result.status is RunStatus.WARN
    assert verifier.calls
    assert control.state.value == "PAUSED"
    assert any("fidelity" in reason.lower() or "table count mismatch" in reason.lower() for reason in result.reasons)
    # The executor must return at the failing safe boundary, before the next event is emitted.
    end_table_index = next(i for i,e in enumerate(prepared.blueprint.events) if e.event_type == "EndTable")
    assert len(controller.events) == end_table_index + 1


def test_prepare_builds_interactive_metadata_policy_from_source(tmp_path):
    from docx import Document
    from word_replica.domain.enums import MetadataMode

    source = tmp_path / "metadata.docx"
    doc = Document(); doc.add_paragraph("Body")
    doc.core_properties.title = "Research title"
    doc.core_properties.subject = "Policy analysis"
    doc.save(source)
    store = ProjectStore(tmp_path / "projects")
    options = RebuildOptions(
        reconstruction_mode=ReconstructionMode.INTERACTIVE,
        metadata=MetadataMode.PRESERVE,
    )
    paths = store.create_project(source, options)
    prepared = InteractiveRebuildService(word_probe=lambda: True).prepare(
        source, options, paths, AuditLog(paths.logs_dir / "audit.jsonl")
    )
    assert prepared.model.extras["metadata_policy"]["title"] == "Research title"
    assert prepared.model.extras["metadata_policy"]["subject"] == "Policy analysis"


def test_resume_checkpoint_preserves_metadata_mode_settings(tmp_path):
    from word_replica.domain.enums import MetadataMode
    service = InteractiveRebuildService()
    options = RebuildOptions(
        reconstruction_mode=ReconstructionMode.INTERACTIVE,
        metadata=MetadataMode.PRESERVE,
        preserve_author_fields=True,
        custom_metadata_allowlist=("StudyId",),
    )
    settings = service._run_settings_dict(options)
    restored = service._rebuild_options_from_settings(settings)
    assert restored.metadata is MetadataMode.PRESERVE
    assert restored.preserve_author_fields is True
    assert restored.custom_metadata_allowlist == ("StudyId",)


def test_prepare_honors_disabled_table_fast_path(tmp_path):
    from docx import Document

    source = tmp_path / "legacy-table.docx"
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Alpha"
    document.save(source)
    options = RebuildOptions(
        reconstruction_mode=ReconstructionMode.INTERACTIVE,
        interactive=InteractiveOptions(enable_table_fast_path=False),
    )
    store = ProjectStore(tmp_path / "projects")
    paths = store.create_project(source, options)

    prepared = InteractiveRebuildService(word_probe=lambda: True).prepare(
        source,
        options,
        paths,
        AuditLog(paths.logs_dir / "audit.jsonl"),
    )

    assert not any(
        event.event_type == "InsertTableBatch"
        for event in prepared.blueprint.events
    )
    assert any(
        event.event_type == "BeginTable"
        for event in prepared.blueprint.events
    )


def test_prepare_and_resume_recompute_identical_model_fingerprint_with_metadata_policy(tmp_path):
    from docx import Document
    from word_replica.domain.enums import MetadataMode
    from word_replica.interactive.blueprint import BlueprintCompiler

    source = tmp_path / "resume_metadata.docx"
    doc = Document(); doc.add_paragraph("Body"); doc.core_properties.title = "Stable title"; doc.save(source)
    store = ProjectStore(tmp_path / "projects")
    options = RebuildOptions(
        reconstruction_mode=ReconstructionMode.INTERACTIVE,
        metadata=MetadataMode.PRESERVE,
    )
    paths = store.create_project(source, options)
    service = InteractiveRebuildService(word_probe=lambda: True)
    prepared = service.prepare(source, options, paths, AuditLog(paths.logs_dir / "audit.jsonl"))
    restored_options = service._rebuild_options_from_settings(service._run_settings_dict(options))
    restored_model = service._parse_model_for_options(source, restored_options)
    restored_blueprint = BlueprintCompiler().compile(restored_model)
    assert restored_blueprint.source_model_fingerprint == prepared.blueprint.source_model_fingerprint
    assert restored_blueprint.fingerprint == prepared.blueprint.fingerprint
