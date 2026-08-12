from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK


def build_core_fixture(path: Path) -> Path:
    doc = Document()
    heading = doc.add_heading("Heading One", level=1)
    heading.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.add_run("Bold").bold = True
    p.add_run().add_tab()
    p.add_run(" normal")
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("after break")
    section = doc.add_section()
    section.orientation = WD_ORIENT.LANDSCAPE
    doc.save(path)
    return path

from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo
from lxml import etree
from PIL import Image
from docx.shared import Inches

_CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def build_extended_fixture(path: Path) -> Path:
    png = path.with_suffix(".png")
    Image.new("RGB", (8, 8), "white").save(png)
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "merged"
    doc.add_paragraph().add_run().add_picture(str(png), width=Inches(0.25))
    doc.sections[0].header.paragraphs[0].text = "Header"
    doc.sections[0].footer.paragraphs[0].text = "Footer"
    doc.save(path)
    patch_notes(path, footnote_text="Footnote text", endnote_text="Endnote text")
    png.unlink()
    return path


def patch_notes(path: Path, footnote_text: str, endnote_text: str) -> None:
    tmp = path.with_suffix(".tmp.docx")
    with ZipFile(path, "r") as src:
        members = {name: src.read(name) for name in src.namelist()}

    ct = etree.fromstring(members["[Content_Types].xml"])
    for part_name, content_type in (
        ("/word/footnotes.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"),
        ("/word/endnotes.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"),
    ):
        if not ct.xpath(f"//*[local-name()='Override'][@PartName='{part_name}']"):
            node = etree.SubElement(ct, f"{{{_CT_NS}}}Override")
            node.set("PartName", part_name)
            node.set("ContentType", content_type)
    members["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone="yes")

    rel_name = "word/_rels/document.xml.rels"
    rels = etree.fromstring(members[rel_name])
    existing_ids = {node.get("Id") for node in rels}
    foot_rid = "rIdFootnote"
    end_rid = "rIdEndnote"
    while foot_rid in existing_ids:
        foot_rid += "X"
    while end_rid in existing_ids or end_rid == foot_rid:
        end_rid += "X"
    for rid, target, rel_type in (
        (foot_rid, "footnotes.xml", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"),
        (end_rid, "endnotes.xml", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes"),
    ):
        node = etree.SubElement(rels, f"{{{_REL_NS}}}Relationship")
        node.set("Id", rid)
        node.set("Type", rel_type)
        node.set("Target", target)
    members[rel_name] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone="yes")

    def notes_xml(kind: str, text: str) -> bytes:
        root = etree.Element(f"{{{_W_NS}}}{kind}s", nsmap={"w": _W_NS})
        note = etree.SubElement(root, f"{{{_W_NS}}}{kind}")
        note.set(f"{{{_W_NS}}}id", "1")
        p = etree.SubElement(note, f"{{{_W_NS}}}p")
        r = etree.SubElement(p, f"{{{_W_NS}}}r")
        t = etree.SubElement(r, f"{{{_W_NS}}}t")
        t.text = text
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")

    members["word/footnotes.xml"] = notes_xml("footnote", footnote_text)
    members["word/endnotes.xml"] = notes_xml("endnote", endnote_text)

    document = etree.fromstring(members["word/document.xml"])
    body = document.find(f"{{{_W_NS}}}body")
    first_p = body.find(f"{{{_W_NS}}}p") if body is not None else None
    if first_p is None and body is not None:
        first_p = etree.SubElement(body, f"{{{_W_NS}}}p")
    if first_p is not None:
        for tag in ("footnoteReference", "endnoteReference"):
            r = etree.SubElement(first_p, f"{{{_W_NS}}}r")
            ref = etree.SubElement(r, f"{{{_W_NS}}}{tag}")
            ref.set(f"{{{_W_NS}}}id", "1")
    members["word/document.xml"] = etree.tostring(document, xml_declaration=True, encoding="UTF-8", standalone="yes")

    with ZipFile(tmp, "w", ZIP_DEFLATED) as dst:
        for name, data in members.items():
            dst.writestr(name, data)
    tmp.replace(path)


def build_review_fixture(path: Path) -> Path:
    doc = Document()
    p = doc.add_paragraph("base ")
    hidden = p.add_run("hidden text")
    hidden.font.hidden = True
    doc.save(path)
    _patch_review_structures(path)
    return path


def _patch_review_structures(path: Path) -> None:
    tmp = path.with_suffix(".review.tmp.docx")
    with ZipFile(path, "r") as src:
        members = {name: src.read(name) for name in src.namelist()}

    document = etree.fromstring(members["word/document.xml"])
    body = document.find(f"{{{_W_NS}}}body")
    p = body.find(f"{{{_W_NS}}}p")

    bookmark_start = etree.Element(f"{{{_W_NS}}}bookmarkStart")
    bookmark_start.set(f"{{{_W_NS}}}id", "7")
    bookmark_start.set(f"{{{_W_NS}}}name", "EvidenceBookmark")
    p.insert(0, bookmark_start)
    bookmark_end = etree.Element(f"{{{_W_NS}}}bookmarkEnd")
    bookmark_end.set(f"{{{_W_NS}}}id", "7")
    p.append(bookmark_end)

    ins = etree.SubElement(p, f"{{{_W_NS}}}ins")
    ins.set(f"{{{_W_NS}}}id", "10")
    ins.set(f"{{{_W_NS}}}author", "Alice")
    ins.set(f"{{{_W_NS}}}date", "2026-08-11T01:00:00Z")
    ins_r = etree.SubElement(ins, f"{{{_W_NS}}}r")
    ins_t = etree.SubElement(ins_r, f"{{{_W_NS}}}t")
    ins_t.text = "inserted text"

    deleted = etree.SubElement(p, f"{{{_W_NS}}}del")
    deleted.set(f"{{{_W_NS}}}id", "11")
    deleted.set(f"{{{_W_NS}}}author", "Alice")
    deleted.set(f"{{{_W_NS}}}date", "2026-08-11T01:01:00Z")
    del_r = etree.SubElement(deleted, f"{{{_W_NS}}}r")
    del_t = etree.SubElement(del_r, f"{{{_W_NS}}}delText")
    del_t.text = "deleted text"

    field_begin_r = etree.SubElement(p, f"{{{_W_NS}}}r")
    field_begin = etree.SubElement(field_begin_r, f"{{{_W_NS}}}fldChar")
    field_begin.set(f"{{{_W_NS}}}fldCharType", "begin")
    instr_r = etree.SubElement(p, f"{{{_W_NS}}}r")
    instr = etree.SubElement(instr_r, f"{{{_W_NS}}}instrText")
    instr.text = " PAGE "
    sep_r = etree.SubElement(p, f"{{{_W_NS}}}r")
    sep = etree.SubElement(sep_r, f"{{{_W_NS}}}fldChar")
    sep.set(f"{{{_W_NS}}}fldCharType", "separate")
    result_r = etree.SubElement(p, f"{{{_W_NS}}}r")
    result_t = etree.SubElement(result_r, f"{{{_W_NS}}}t")
    result_t.text = "1"
    end_r = etree.SubElement(p, f"{{{_W_NS}}}r")
    end = etree.SubElement(end_r, f"{{{_W_NS}}}fldChar")
    end.set(f"{{{_W_NS}}}fldCharType", "end")

    comment_start = etree.Element(f"{{{_W_NS}}}commentRangeStart")
    comment_start.set(f"{{{_W_NS}}}id", "3")
    p.insert(1, comment_start)
    comment_end = etree.SubElement(p, f"{{{_W_NS}}}commentRangeEnd")
    comment_end.set(f"{{{_W_NS}}}id", "3")
    comment_ref_r = etree.SubElement(p, f"{{{_W_NS}}}r")
    comment_ref = etree.SubElement(comment_ref_r, f"{{{_W_NS}}}commentReference")
    comment_ref.set(f"{{{_W_NS}}}id", "3")
    members["word/document.xml"] = etree.tostring(document, xml_declaration=True, encoding="UTF-8", standalone="yes")

    settings = etree.fromstring(members["word/settings.xml"])
    if settings.find(f"{{{_W_NS}}}trackRevisions") is None:
        settings.insert(0, etree.Element(f"{{{_W_NS}}}trackRevisions"))
    members["word/settings.xml"] = etree.tostring(settings, xml_declaration=True, encoding="UTF-8", standalone="yes")

    comments = etree.Element(f"{{{_W_NS}}}comments", nsmap={"w": _W_NS})
    comment = etree.SubElement(comments, f"{{{_W_NS}}}comment")
    comment.set(f"{{{_W_NS}}}id", "3")
    comment.set(f"{{{_W_NS}}}author", "Reviewer")
    comment.set(f"{{{_W_NS}}}date", "2026-08-11T01:02:00Z")
    cp = etree.SubElement(comment, f"{{{_W_NS}}}p")
    cr = etree.SubElement(cp, f"{{{_W_NS}}}r")
    ct = etree.SubElement(cr, f"{{{_W_NS}}}t")
    ct.text = "Review comment"
    members["word/comments.xml"] = etree.tostring(comments, xml_declaration=True, encoding="UTF-8", standalone="yes")

    content_types = etree.fromstring(members["[Content_Types].xml"])
    if not content_types.xpath("//*[local-name()='Override'][@PartName='/word/comments.xml']"):
        override = etree.SubElement(content_types, f"{{{_CT_NS}}}Override")
        override.set("PartName", "/word/comments.xml")
        override.set("ContentType", "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml")
    members["[Content_Types].xml"] = etree.tostring(content_types, xml_declaration=True, encoding="UTF-8", standalone="yes")

    rels = etree.fromstring(members["word/_rels/document.xml.rels"])
    rel = etree.SubElement(rels, f"{{{_REL_NS}}}Relationship")
    rel.set("Id", "rIdCommentsReplica")
    rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments")
    rel.set("Target", "comments.xml")
    members["word/_rels/document.xml.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone="yes")

    with ZipFile(tmp, "w", ZIP_DEFLATED) as dst:
        for name, data in members.items():
            dst.writestr(name, data)
    tmp.replace(path)

# --- v1 deterministic acceptance corpus -------------------------------------------------
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

_FIXED_W3CDTF = "2026-08-11T00:00:00Z"


def _atomic_patch(path: Path, mutator) -> None:
    tmp = path.with_suffix(path.suffix + ".patch.tmp")
    with ZipFile(path, "r") as src:
        members = {name: src.read(name) for name in src.namelist()}
    mutator(members)
    with ZipFile(tmp, "w", ZIP_DEFLATED) as dst:
        for name in sorted(members):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dst.writestr(info, members[name])
    with ZipFile(tmp, "r") as check:
        assert check.testzip() is None
    tmp.replace(path)


def _normalize_fixture(path: Path) -> None:
    def mutate(members: dict[str, bytes]) -> None:
        core_name = "docProps/core.xml"
        if core_name in members:
            root = etree.fromstring(members[core_name])
            ns = {
                "dcterms": "http://purl.org/dc/terms/",
                "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
            }
            for tag in ("created", "modified"):
                node = root.find(f"dcterms:{tag}", namespaces=ns)
                if node is not None:
                    node.text = _FIXED_W3CDTF
            revision = root.find("cp:revision", namespaces=ns)
            if revision is not None:
                revision.text = "77"
            members[core_name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
        app_name = "docProps/app.xml"
        if app_name in members:
            app = etree.fromstring(members[app_name])
            ep_ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
            total = app.find(f"{{{ep_ns}}}TotalTime")
            if total is not None:
                total.text = "999"
            members[app_name] = etree.tostring(app, xml_declaration=True, encoding="UTF-8", standalone="yes")
    _atomic_patch(path, mutate)


def build_plain_text(path: Path) -> Path:
    doc = Document(); doc.add_paragraph("Plain paragraph one."); doc.add_paragraph("Plain paragraph two."); doc.save(path); return path


def build_headings_styles(path: Path) -> Path:
    doc = Document(); doc.add_heading("Research Title", 1); doc.add_heading("Subsection", 2)
    style = doc.styles.add_style("ReplicaBody", WD_STYLE_TYPE.PARAGRAPH); style.font.name = "Arial"; style.font.size = Pt(11)
    doc.add_paragraph("Styled body content.", style="ReplicaBody"); doc.save(path); return path


def build_lists(path: Path) -> Path:
    doc = Document()
    for text in ("First", "Second", "Third"): doc.add_paragraph(text, style="List Number")
    for text in ("Alpha", "Beta"): doc.add_paragraph(text, style="List Bullet")
    doc.save(path); return path


def build_tables_merged(path: Path) -> Path:
    doc = Document(); table = doc.add_table(rows=3, cols=3); table.style = "Table Grid"
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Vertical"
    table.cell(1, 1).text = "A"; table.cell(1, 2).text = "B"; table.cell(2, 1).text = "C"; table.cell(2, 2).text = "D"
    doc.save(path); return path


def build_images_inline_floating(path: Path) -> Path:
    image_path = path.with_suffix(".fixture.png")
    Image.new("RGB", (40, 24), "white").save(image_path)
    doc = Document(); p = doc.add_paragraph("Image: "); p.add_run().add_picture(str(image_path), width=Inches(0.6)); doc.add_paragraph("After image."); doc.save(path)
    image_path.unlink(missing_ok=True); return path


def build_sections_orientations(path: Path) -> Path:
    doc = Document(); doc.add_paragraph("Portrait section")
    landscape = doc.add_section(); landscape.orientation = WD_ORIENT.LANDSCAPE; landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    doc.add_paragraph("Landscape section"); doc.add_section(); doc.add_paragraph("Portrait again"); doc.save(path); return path


def _append_field_to_paragraph(p_node, instruction: str, result: str) -> None:
    for kind in ("begin",):
        r = etree.SubElement(p_node, f"{{{_W_NS}}}r"); fld = etree.SubElement(r, f"{{{_W_NS}}}fldChar"); fld.set(f"{{{_W_NS}}}fldCharType", kind)
    r = etree.SubElement(p_node, f"{{{_W_NS}}}r"); instr = etree.SubElement(r, f"{{{_W_NS}}}instrText"); instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); instr.text = f" {instruction} "
    r = etree.SubElement(p_node, f"{{{_W_NS}}}r"); fld = etree.SubElement(r, f"{{{_W_NS}}}fldChar"); fld.set(f"{{{_W_NS}}}fldCharType", "separate")
    r = etree.SubElement(p_node, f"{{{_W_NS}}}r"); t = etree.SubElement(r, f"{{{_W_NS}}}t"); t.text = result
    r = etree.SubElement(p_node, f"{{{_W_NS}}}r"); fld = etree.SubElement(r, f"{{{_W_NS}}}fldChar"); fld.set(f"{{{_W_NS}}}fldCharType", "end")


def build_headers_footers_numbers(path: Path) -> Path:
    doc = Document(); doc.add_paragraph("Body with header/footer")
    sec = doc.sections[0]; sec.header.paragraphs[0].text = "Institution Header"; sec.footer.paragraphs[0].text = "Page "
    doc.save(path)
    def mutate(members):
        footer = etree.fromstring(members["word/footer1.xml"]); p = footer.find(f"{{{_W_NS}}}p"); _append_field_to_paragraph(p, "PAGE", "1")
        members["word/footer1.xml"] = etree.tostring(footer, xml_declaration=True, encoding="UTF-8", standalone="yes")
    _atomic_patch(path, mutate); return path


def build_footnotes_endnotes(path: Path) -> Path:
    doc = Document(); doc.add_paragraph("Statement with notes"); doc.save(path); patch_notes(path, "Footnote evidence", "Endnote evidence"); return path


def build_toc_fields(path: Path) -> Path:
    doc = Document(); doc.add_heading("Contents", 1); doc.add_paragraph("TOC placeholder"); doc.add_heading("Chapter One", 1); doc.add_paragraph("Chapter body"); doc.save(path)
    def mutate(members):
        root = etree.fromstring(members["word/document.xml"]); p = root.xpath("//w:p", namespaces={"w": _W_NS})[1]; _append_field_to_paragraph(p, 'TOC \\o "1-3" \\h \\z \\u', "Chapter One .... 1")
        members["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
    _atomic_patch(path, mutate); return path


def build_comments_tracked_changes(path: Path) -> Path:
    return build_review_fixture(path)


def build_bookmarks_crossrefs(path: Path) -> Path:
    doc = Document(); doc.add_paragraph("Target paragraph"); doc.add_paragraph("See target: "); doc.save(path)
    def mutate(members):
        root = etree.fromstring(members["word/document.xml"]); ps = root.xpath("//w:body/w:p", namespaces={"w": _W_NS})
        start = etree.Element(f"{{{_W_NS}}}bookmarkStart"); start.set(f"{{{_W_NS}}}id", "5"); start.set(f"{{{_W_NS}}}name", "TargetBookmark"); ps[0].insert(0, start)
        end = etree.SubElement(ps[0], f"{{{_W_NS}}}bookmarkEnd"); end.set(f"{{{_W_NS}}}id", "5")
        _append_field_to_paragraph(ps[1], "REF TargetBookmark \\h", "Target paragraph")
        members["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
    _atomic_patch(path, mutate); return path


def _patch_embedded_parts(path: Path) -> None:
    def mutate(members):
        members["word/embeddings/object1.bin"] = b"WORD_REPLICA_DETERMINISTIC_EMBEDDED_OBJECT"
        chart = etree.Element("{http://schemas.openxmlformats.org/drawingml/2006/chart}chartSpace", nsmap={"c":"http://schemas.openxmlformats.org/drawingml/2006/chart"})
        members["word/charts/chart1.xml"] = etree.tostring(chart, xml_declaration=True, encoding="UTF-8", standalone="yes")
        ct = etree.fromstring(members["[Content_Types].xml"])
        for part_name, content_type in (("/word/charts/chart1.xml", "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"), ("/word/embeddings/object1.bin", "application/vnd.openxmlformats-officedocument.oleObject")):
            if not ct.xpath(f"//*[local-name()='Override'][@PartName='{part_name}']"):
                o = etree.SubElement(ct, f"{{{_CT_NS}}}Override"); o.set("PartName", part_name); o.set("ContentType", content_type)
        members["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone="yes")
    _atomic_patch(path, mutate)


def build_charts_embedded(path: Path) -> Path:
    doc = Document(); doc.add_paragraph("Embedded chart/object evidence"); doc.save(path); _patch_embedded_parts(path); return path


def build_academic_complex(path: Path) -> Path:
    image_path = path.with_suffix(".academic.png"); Image.new("RGB", (64, 40), "white").save(image_path)
    doc = Document(); doc.add_heading("Academic Research", 1); doc.add_paragraph("Abstract with a reproducible statement.")
    doc.add_heading("Results", 2); table = doc.add_table(rows=2, cols=2); table.cell(0,0).text="Measure"; table.cell(0,1).text="Value"; table.cell(1,0).text="A"; table.cell(1,1).text="42"
    doc.add_paragraph().add_run().add_picture(str(image_path), width=Inches(0.7)); doc.sections[0].header.paragraphs[0].text="Academic Header"; doc.sections[0].footer.paragraphs[0].text="Academic Footer"; doc.save(path)
    image_path.unlink(missing_ok=True); patch_notes(path, "Academic footnote", "Academic endnote"); _patch_review_structures(path); _patch_embedded_parts(path); return path


def build_interactive_mixed_breaks(path: Path) -> Path:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Normal ")
    bold = p.add_run("Bold"); bold.bold = True
    p.add_run().add_tab()
    italic = p.add_run("Italic"); italic.italic = True
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run("Line two")
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("After page break")
    doc.save(path)
    return path


def build_nested_table(path: Path) -> Path:
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    outer.style = "Table Grid"
    outer.cell(0, 0).paragraphs[0].add_run("Outer start")
    nested = outer.cell(0, 0).add_table(rows=2, cols=2)
    nested.style = "Table Grid"
    nested.cell(0, 0).text = "N11"
    nested.cell(0, 1).text = "N12"
    nested.cell(1, 0).text = "N21"
    nested.cell(1, 1).text = "N22"
    outer.cell(0, 0).add_paragraph("Outer end")
    doc.save(path)
    return path


def build_table_shading_borders(path: Path) -> Path:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for r in range(2):
        for c in range(2):
            cell = table.cell(r, c)
            cell.text = f"R{r+1}C{c+1}"
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "D9EAF7" if r == 0 else "FFFFFF")
            tc_pr.append(shd)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Shaded merged heading"
    doc.save(path)
    return path


def build_academic_supported(path: Path) -> Path:
    image_path = path.with_suffix(".supported.png")
    Image.new("RGB", (64, 40), "white").save(image_path)
    doc = Document()
    doc.add_heading("Supported Academic", 1)
    doc.add_paragraph("Abstract statement with mixed ").add_run("emphasis").bold = True
    table = doc.add_table(rows=2, cols=2); table.style = "Table Grid"
    table.cell(0,0).text="Measure"; table.cell(0,1).text="Value"; table.cell(1,0).text="A"; table.cell(1,1).text="42"
    doc.add_paragraph().add_run().add_picture(str(image_path), width=Inches(0.7))
    doc.sections[0].header.paragraphs[0].text="Supported Header"
    doc.sections[0].footer.paragraphs[0].text="Supported Footer"
    doc.save(path)
    image_path.unlink(missing_ok=True)
    patch_notes(path, "Supported footnote", "Supported endnote")
    return path


BUILDERS = [
    ("01_plain_text.docx", build_plain_text),
    ("02_headings_styles.docx", build_headings_styles),
    ("03_lists.docx", build_lists),
    ("04_tables_merged.docx", build_tables_merged),
    ("05_images_inline_floating.docx", build_images_inline_floating),
    ("06_sections_orientations.docx", build_sections_orientations),
    ("07_headers_footers_numbers.docx", build_headers_footers_numbers),
    ("08_footnotes_endnotes.docx", build_footnotes_endnotes),
    ("09_toc_fields.docx", build_toc_fields),
    ("10_comments_tracked_changes.docx", build_comments_tracked_changes),
    ("11_bookmarks_crossrefs.docx", build_bookmarks_crossrefs),
    ("12_charts_embedded.docx", build_charts_embedded),
    ("13_academic_complex.docx", build_academic_complex),
    ("14_interactive_mixed_breaks.docx", build_interactive_mixed_breaks),
    ("15_nested_table.docx", build_nested_table),
    ("16_table_shading_borders.docx", build_table_shading_borders),
    ("17_academic_supported.docx", build_academic_supported),
]


def build_all_corpus(out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    built: list[Path] = []
    for filename, builder in BUILDERS:
        target = out_dir / filename
        builder(target)
        _normalize_fixture(target)
        with ZipFile(target) as archive:
            assert archive.testzip() is None
        built.append(target)
    return built
